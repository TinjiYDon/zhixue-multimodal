# -*- coding: utf-8 -*-
"""Generate two zhixue closing-report Word documents."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

OUT = Path(r"d:\project\zhixue-multimodal\docs")


def set_run_font(run, name="宋体", size=12, bold=False, east_asia=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia or name)


def add_heading_cn(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, "黑体", 16 if level == 1 else 14 if level == 2 else 12, bold=True)
    return p


def add_para(doc, text, first_line_indent=True, size=12, bold=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(6)
    if first_line_indent:
        pf.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    set_run_font(run, "宋体", size, bold=bold)
    return p


def add_bullet(doc, text, size=12):
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    run = p.add_run(text)
    set_run_font(run, "宋体", size)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    return p


def setup_doc(title: str) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.17)
    section.right_margin = Cm(3.17)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(title)
    set_run_font(r, "黑体", 18, bold=True)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("智学多模态 Agent（zhixue-multimodal）")
    set_run_font(sr, "楷体", 12)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mr = meta.add_run("仓库：TinjiYDon/zhixue-multimodal　·　文稿日期：2026-08-14")
    set_run_font(mr, "宋体", 10)
    return doc


def build_part1():
    doc = setup_doc("结项报告 · 第 1 部分\n模型架构与算法（可用正文稿）")
    note = doc.add_paragraph()
    nr = note.add_run(
        "说明：本稿按结项报告书语气撰写，可直接删改后作为「第 1 部分」正文。"
        "表述与当前工程实现一致：正式 RAG/真 WhisperX 为演进项，勿改写成已全面生产上线。"
    )
    set_run_font(nr, "楷体", 10)
    note.paragraph_format.space_after = Pt(12)

    add_heading_cn(doc, "1.1 问题定义与目标", 1)
    add_para(
        doc,
        "课堂场景中，音视频、板书/PPT 与师生问答往往割裂存储：学生复习时难以按时间点回看讲解，"
        "教师也难以基于整节课内容做即时答疑。智学多模态 Agent 面向「课程级」多模态理解，目标是："
        "在一次课程媒体上传后，自动完成语音转写与版面文字提取，形成带时间戳的课程时间轴（timeline），"
        "并在此基础上提供基于检索增强（RAG）的课程问答能力，支撑 Web 端与微信小程序端的学习交互。",
    )
    add_para(
        doc,
        "本项目强调工程可复现与契约先行：转写与 OCR 输出统一到 TranscriptResult / OcrResult 等模式；"
        "无 GPU 或未安装重型模型时，可通过 fixture 后端完成联调与验收，避免将「演示可跑」等同于"
        "「生产级语音模型已默认上线」。本系统与 ICU 临床预警/床位调度项目零代码耦合，仅在工程方法上可并列对照。",
    )

    add_heading_cn(doc, "1.2 总体算法与数据流", 1)
    add_para(
        doc,
        "端到端主链路可概括为：媒体对象入库 → 异步任务调度 → 音视频预处理与多模态感知 → "
        "页级/时间对齐 → 时间轴物化 → 问答检索与生成。形式化地，对课程 c，设媒体对象为 M_c，"
        "系统输出时间轴 T_c = {(t_i^s, t_i^e, text_i)} 与可选幻灯片页序列 S_c，"
        "问答接口在给定问题 q 时返回答案 a 及证据片段集合 E ⊆ T_c ∪ S_c。",
    )
    add_para(doc, "关键阶段如下。", first_line_indent=False)
    add_bullet(doc, "上传与对象存储：客户端经预签名 URL 将媒体写入 MinIO；complete 校验对象存在后创建 Job。")
    add_bullet(doc, "Worker 调度：BackgroundTasks（后续可换队列）调用多媒体流水线，更新 Job 进度与状态。")
    add_bullet(doc, "感知：FFmpeg 抽取 16 kHz 单声道 WAV；ASR 生成带时间戳字幕；OCR 提取板书/PPT 文本与框。")
    add_bullet(doc, "对齐：将字幕片段与 OCR 页在时间或语义上对齐，形成统一 timeline（当前含工具函数与占位管线）。")
    add_bullet(doc, "问答：基于课程维度隔离的检索（现为内存/占位 RAG；规划 pgvector + LLM）返回答案与来源。")

    add_heading_cn(doc, "1.3 多媒体感知模块", 1)
    add_heading_cn(doc, "1.3.1 音频预处理", 2)
    add_para(
        doc,
        "为降低 ASR 对采样率与声道的敏感性，系统使用 FFmpeg 将输入媒体解封装并重采样为 16 kHz 单声道 PCM WAV。"
        "该步骤与具体容器格式解耦：联调默认推荐 MP4（video/mp4），但上传 API 层当前未强制 MIME 白名单；"
        "实际能否抽音取决于本机 FFmpeg 对容器的支持能力。预处理失败将直接导致转写任务失败，因而部署环境需保证 ffmpeg 在 PATH 中可用。",
    )

    add_heading_cn(doc, "1.3.2 语音识别与时间对齐（ASR）", 2)
    add_para(
        doc,
        "语音识别路径以 Whisper 族模型为核心，工程上支持 WhisperX / faster-whisper 等实现，并可选字词级强制对齐，"
        "以获得更细粒度的 (text, start, end) 片段。输出严格映射到团队约定的 TranscriptResult / TranscriptSegment 契约，"
        "便于后续对齐、入库与前端时间轴渲染。针对不同 Python/依赖版本的参数差异，实现侧引入兼容补丁，保证调用签名稳定。",
    )
    add_para(
        doc,
        "在无 GPU、未安装重型依赖或 CI 场景下，系统可通过环境变量切换至 fixture 后端，返回符合 schema 的样例转写结果。"
        "该设计保证「契约测试与联调」与「真模型推理」分离：fixture 通过不等于宣称生产环境已默认启用完整 WhisperX。",
    )

    add_heading_cn(doc, "1.3.3 版面文字识别（OCR）", 2)
    add_para(
        doc,
        "对课件页或板书截图，采用 RapidOCR（ONNXRuntime）等轻量方案提取文本块及四点坐标框，"
        "组装为 OcrResult / OcrPageResult。OCR 与 ASR 并列构成多模态证据：前者偏「看见的板书结构」，"
        "后者偏「听见的讲解时间线」。二者经对齐模块融合后，才能支撑「问一句、跳到对应讲解片段」的交互。",
    )

    add_heading_cn(doc, "1.4 对齐与时间轴构建", 1)
    add_para(
        doc,
        "对齐（alignment）模块负责将离散的 ASR 片段与 OCR 页关联到统一时间轴。当前仓库已提供页级对齐相关工具函数与"
        "timeline API（含 from-fixture 钩子），用于在多媒体与 Job 持久化完备前打通前后端。完整策略可包含："
        "基于时间重叠的启发式匹配、基于编辑距离/语义相似度的字幕—页标题对齐，以及课程维度的隔离存储。"
        "时间轴对外以 GET /api/v1/courses/{id}/timeline 暴露，供 Web 播放器同步字幕、PPT 高亮与进度跳转。",
    )

    add_heading_cn(doc, "1.5 检索增强问答（RAG）", 1)
    add_para(
        doc,
        "问答接口约定为 POST /api/v1/courses/{id}/ask（禁止使用无课程作用域的全局 /ask），以保证多课程数据隔离。"
        "算法流程为：对问题 q 做检索 → 取 Top-K 证据片段 → 拼装提示词 → 调用生成模型得到答案，并回传来源列表。"
        "现阶段实现为可演示的占位 RAG（可命中 fixture timeline 文本）；下一阶段规划引入 pgvector 向量索引与正式 LLM，"
        "在保持同一 API 契约的前提下替换检索与生成后端。评测上应报告检索命中率、答案忠实度与延迟，"
        "并明确区分 fixture 演示与向量库+LLM 正式链路。",
    )

    add_heading_cn(doc, "1.6 与系统边界相关的算法约束", 1)
    add_bullet(doc, "课程隔离：所有检索与时间轴必须以 course_id 为键，禁止跨课串数据。")
    add_bullet(doc, "异步与可观测：感知计算放入 Job/Worker，API 层保持快速返回；状态可查询。")
    add_bullet(doc, "降级优先：重模型不可用时 fixture 保契约，避免阻塞前后端与验收。")
    add_bullet(doc, "诚实表述：不将占位 RAG、模拟视频区或 fixture ASR 写成生产级已交付能力。")

    add_heading_cn(doc, "1.7 小结", 1)
    add_para(
        doc,
        "智学多模态 Agent 的「模型与算法」主线是多媒体感知（ASR/OCR）+ 时间对齐 + 课程级 RAG，"
        "以清晰的数据契约贯通存储、Worker 与前后端。当前版本已具备可演示、可测试的算法骨架与占位智能链路；"
        "正式向量检索与生产级语音识别属于明确规划中的增强项，应在结项材料中按「已实现 / 演进中」分层陈述。",
    )

    path = OUT / "结项报告-第1部分-模型架构与算法.docx"
    doc.save(path)
    return path


def build_intro_234():
    doc = setup_doc("智学多模态 Agent\n项目详细介绍与结项材料（第 2–4 部分）")
    note = doc.add_paragraph()
    nr = note.add_run(
        "结构：前半为项目详细介绍（可直接用于说明书/答辩开篇）；"
        "后半为第 2、3、4 部分关键词与信息补充，供组员扩写。"
    )
    set_run_font(nr, "楷体", 10)
    note.paragraph_format.space_after = Pt(12)

    # —— 详细介绍 ——
    add_heading_cn(doc, "一、项目详细介绍", 1)

    add_heading_cn(doc, "1.1 一句话定位", 2)
    add_para(
        doc,
        "智学多模态 Agent 是面向课堂教学的多模态实时理解系统：完成课程音视频上传后，"
        "自动转写与版面识别，生成可交互时间轴，并支持基于课程内容的智能问答；"
        "提供 Vue3 Web 与 UniApp 微信小程序两端入口。",
    )

    add_heading_cn(doc, "1.2 背景与意义", 2)
    add_para(
        doc,
        "传统网课回放多为「整段视频 + 静态课件」，缺少细粒度时间索引与可追溯答疑。"
        "本项目将 ASR、OCR、时间轴与 RAG 串成一条可工程化交付的链路，使「听到的」与「看到的」"
        "在同一时间坐标系下可检索、可跳转、可引用，从而提升复习效率与答疑质量。"
        "项目坚持契约优先与 fixture 降级，适配教学团队在普通开发机上的联调与答辩演示。",
    )

    add_heading_cn(doc, "1.3 功能范围", 2)
    add_bullet(doc, "课程管理：Course 创建与查询（已落 PostgreSQL）。")
    add_bullet(doc, "媒体上传：MinIO 预签名上传、complete 校验、自动创建转写 Job。")
    add_bullet(doc, "多媒体处理：FFmpeg 抽音、ASR、OCR（真模型可选；fixture 默认可测）。")
    add_bullet(doc, "时间轴：GET timeline；支持 from-fixture 注入演示数据。")
    add_bullet(doc, "问答：POST /api/v1/courses/{id}/ask（课程隔离）。")
    add_bullet(doc, "前端：Web 课程列表/详情（字幕、时间轴、PPT、问答面板）；小程序列表与问答骨架。")

    add_heading_cn(doc, "1.4 技术架构", 2)
    add_para(
        doc,
        "后端采用 FastAPI，前端为 Vue3 + Vite，小程序为 UniApp。基础设施通过 Docker Compose 提供 "
        "PostgreSQL（默认主机端口 5435）、Redis 与 MinIO。对象存储承载媒体文件；关系库承载 Course/Job 等业务实体；"
        "异步任务当前以 BackgroundTasks 触发 Worker，后续可演进为独立队列。API 契约以 OpenAPI（/docs）为准。",
    )
    add_para(doc, "逻辑分层示意：", first_line_indent=False)
    add_para(
        doc,
        "Web / miniapp → FastAPI（upload / courses / jobs / timeline / ask）→ "
        "PostgreSQL + MinIO +（可选）Redis；Worker → multimedia（ffmpeg / ASR / OCR）→ timeline / RAG。",
        first_line_indent=False,
    )

    add_heading_cn(doc, "1.5 协作与里程碑（摘要）", 2)
    add_bullet(doc, "P0 多媒体 schema + fixture：已完成（PR #12）。")
    add_bullet(doc, "MILE-1 Course/Job PostgreSQL：已完成。")
    add_bullet(doc, "MILE-2 Web 接真 timeline/ask：骨架已完成（Issue #7 仍 open，待关验收标准）。")
    add_bullet(doc, "MILE-3 小程序骨架：已完成（PR #11）。")
    add_bullet(doc, "正式 pgvector RAG、MILE-4 答辩彩排：进行中/未做。")
    add_bullet(doc, "人工验收（2026-08-14）：有条件通过；详见 docs/ACCEPTANCE_ZHIXUE.md。")

    add_heading_cn(doc, "1.6 使用与演示要点", 2)
    add_para(
        doc,
        "本地：docker compose up -d 后启动 uvicorn（8000）与 npm run dev（5173）。"
        "上传联调默认使用 mp4。演示路径建议：健康检查 → 课程列表 → 进入课程 → 刷新/注入 timeline → 提问。"
        "系统当前无账号登录体系，答辩表述需与此一致。Release 标签：zhixue-mile2。",
    )

    add_heading_cn(doc, "1.7 边界与非目标", 2)
    add_bullet(doc, "不与 ICU 预警/床位系统做运行时数据耦合。")
    add_bullet(doc, "不宣称生产级 WhisperX / 正式向量 RAG 已默认上线（除非后续文档改写并验收）。")
    add_bullet(doc, "视频区可为模拟播放器占位；真实流媒体播放为增强项。")

    # —— 第 2 部分 ——
    add_heading_cn(doc, "二、第 2 部分 · 系统实现与工程交付（关键词与信息补充）", 1)
    add_heading_cn(doc, "2.1 建议标题", 2)
    add_para(doc, "系统架构、工程实现与可复现交付", first_line_indent=False)

    add_heading_cn(doc, "2.2 核心关键词", 2)
    for k in [
        "仓独立 / 契约优先 / OpenAPI",
        "FastAPI · Vue3 · UniApp",
        "Docker Compose：PostgreSQL · Redis · MinIO",
        "upload：presign → PUT → complete → 自动 Job",
        "Worker / BackgroundTasks · media_key",
        "FFmpeg 16kHz mono · ASR · OCR · schema 对齐",
        "ZHIXUE_ASR_BACKEND=fixture 降级",
        "timeline / from-fixture · ask 课程隔离",
        "pytest 契约测试 · Release zhixue-mile2",
    ]:
        add_bullet(doc, k)

    add_heading_cn(doc, "2.3 信息补充（可写入正文的事实）", 2)
    add_bullet(doc, "目录：backend/（API+服务+Worker）、web/、miniapp/、docker-compose.yml、docs/。")
    add_bullet(doc, "上传样例：lecture.mp4 / video/mp4；API 未做扩展名强制校验。")
    add_bullet(doc, "转写契约：backend/app/schemas/transcript.py（TranscriptResult / OcrResult）。")
    add_bullet(doc, "ask 路径必须为 POST /api/v1/courses/{id}/ask。")
    add_bullet(doc, "协作文档：TEAM_ASSIGNMENT.md、COLLABORATION.md、BACKLOG.md、OWNER_VS_TEAM.md。")
    add_bullet(doc, "质量门禁：backend pytest；禁止把 node_modules/.venv/大媒体样例无 Git。")
    add_bullet(doc, "已知限制：无登录；视频可为模拟区；RAG 仍可为占位；Issue #6/#7、PR #13 待收尾。")

    add_heading_cn(doc, "2.4 建议章节提纲", 2)
    for i, t in enumerate(
        [
            "总体架构图与模块职责（A Web / B miniapp / C 多媒体 / D Job / 负责人 API+RAG）",
            "上传与对象存储、Job 生命周期",
            "多媒体流水线与 fixture 策略",
            "timeline 与前端交互（字幕、PPT、跳转）",
            "ask 接口与课程隔离",
            "部署、配置与复现步骤",
        ],
        1,
    ):
        add_bullet(doc, f"{i}. {t}")

    # —— 第 3 部分 ——
    add_heading_cn(doc, "三、第 3 部分 · 实验评测与验收（关键词与信息补充）", 1)
    add_heading_cn(doc, "3.1 建议标题", 2)
    add_para(doc, "实验设计、指标与验收结论", first_line_indent=False)

    add_heading_cn(doc, "3.2 核心关键词", 2)
    for k in [
        "契约测试 / schema 对齐",
        "pytest：upload · ask · timeline · jobs · multimedia",
        "冒烟：health → courses → timeline → ask",
        "人工视觉验收 · 有条件通过",
        "fixture ≠ 真 Whisper 上线",
        "course_id 隔离",
        "Open Issue #6 #7 · Open PR #13",
    ]:
        add_bullet(doc, k)

    add_heading_cn(doc, "3.3 信息补充（可引用数据）", 2)
    add_bullet(doc, "2026-08-14：pytest 14 passed。")
    add_bullet(doc, "验收结论：有条件通过（见 ACCEPTANCE_ZHIXUE.md）。")
    add_bullet(doc, "截图：docs/acceptance-screens/01-home.png、02-course.png。")
    add_bullet(doc, "冒烟 API 均为 HTTP 200；后端日志无 Traceback/5xx（该次验收）。")
    add_bullet(doc, "轻微失败：首页「课程列表」标题重复；预期失败：模拟视频区、占位 RAG。")
    add_bullet(doc, "无登录/退出：冒烟用「进入课程→提问→回首页」替代。")

    add_heading_cn(doc, "3.4 建议章节提纲", 2)
    for i, t in enumerate(
        [
            "测试分层：单元契约 / 接口冒烟 / 人工视觉",
            "多媒体 schema 与 fixture 策略说明",
            "验收环境、步骤与结果表",
            "缺陷、风险与边界案例",
            "与下一阶段正式 RAG/真 ASR 的评测计划（可一笔带过）",
        ],
        1,
    ):
        add_bullet(doc, f"{i}. {t}")

    # —— 第 4 部分 ——
    add_heading_cn(doc, "四、第 4 部分 · 总结与展望（关键词与信息补充）", 1)
    add_heading_cn(doc, "4.1 建议标题", 2)
    add_para(doc, "工作总结、创新点、不足与下一步", first_line_indent=False)

    add_heading_cn(doc, "4.2 核心关键词", 2)
    for k in [
        "课堂多模态实时 Agent",
        "时间轴驱动的教与学",
        "契约先行 · fixture 可演示",
        "Web + 小程序双端",
        "pgvector 正式 RAG（展望）",
        "MILE-4 答辩彩排",
        "工程方法可与其他 Agent 项目并列对照（零耦合）",
    ]:
        add_bullet(doc, k)

    add_heading_cn(doc, "4.3 信息补充", 2)
    add_para(doc, "已完成（可写进总结）：", first_line_indent=False)
    add_bullet(doc, "基建与 API 骨架、多媒体契约与 fixture、Course/Job PG、Web 真 API 骨架、小程序 P0、有条件验收通过。")
    add_para(doc, "创新点表述建议：", first_line_indent=False)
    add_bullet(doc, "把 ASR/OCR/时间轴/问答落成可分工、可测试、可演示的垂直切片工程，而不是单次脚本。")
    add_bullet(doc, "用 fixture 分离「联调可用性」与「重模型依赖」，降低五人协作摩擦。")
    add_para(doc, "不足：", first_line_indent=False)
    add_bullet(doc, "正式向量 RAG 与生产 ASR 未默认交付；Issue #6/#7 未关闭；PR #13 待审；无账号体系；UI 仍有小瑕疵。")
    add_para(doc, "下一步（对齐 ROADMAP）：", first_line_indent=False)
    add_bullet(doc, "P0：关或重写 #6/#7；审 PR #13。")
    add_bullet(doc, "P1：pgvector RAG；MILE-4 五分钟彩排。")
    add_bullet(doc, "P2：真 WhisperX 本机验证；可选上传 MIME 白名单。")

    add_heading_cn(doc, "4.4 禁止表述（答辩红线）", 2)
    add_bullet(doc, "「已具备完整账号登录与权限体系」。")
    add_bullet(doc, "「生产级 WhisperX 已默认上线」。")
    add_bullet(doc, "「与 ICU 风险/床位系统已打通」。")

    add_heading_cn(doc, "附录 · 文档与仓库索引", 1)
    add_bullet(doc, "GitHub：https://github.com/TinjiYDon/zhixue-multimodal")
    add_bullet(doc, "进度：docs/PROGRESS.md · 路线：docs/ROADMAP.md · 变更：docs/CHANGELOG.md")
    add_bullet(doc, "验收：docs/ACCEPTANCE_ZHIXUE.md · 本仓结项框架：docs/REPORT_OUTLINE.md")

    path = OUT / "项目详细介绍与结项第2-4部分补充.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    p1 = build_part1()
    p2 = build_intro_234()
    print(p1)
    print(p2)
